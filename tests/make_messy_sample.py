# -*- coding: utf-8 -*-
"""
書式の乱れた就業規則サンプル（tests/samples/messy_kisoku.docx）を作る。

厚労省モデル就業規則（01_01_202601.docx）の第１〜３章を元に、
実在の出来の悪い規則でよく見る乱れを規則的に混ぜる。
  - 番号・見出しに半角／全角スペースが混ざる（「第 3 条」「（ 目 的 ）」）
  - 条見出しと第１文が同じ行、または見出しが条番号の後ろにある
  - 全角数字と半角数字、全角括弧と半角括弧の混在
  - 項番号・号番号の書き方がばらばら（「２」「2.」「2 .」、「①」「(1)」「1)」）
  - 字下げを全角スペースで作る、行末の余分な空白
  - 見た目の改行（文の途中で段落を分ける／段落内改行）
  - 空段落の連続

  python3 tests/make_messy_sample.py
乱れ方は固定（乱数を使わない）なので、何度作っても同じ内容になる。
"""

import os
import re

from docx import Document

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "tests", "samples", "01_01_202601.docx")
DST = os.path.join(ROOT, "tests", "samples", "messy_kisoku.docx")

Z2H = str.maketrans("０１２３４５６７８９（）", "0123456789()")
CIRCLED = "①②③④⑤⑥⑦⑧⑨⑩"


def half(s):
    return s.translate(Z2H)


def main():
    src = [p.text for p in Document(SRC).paragraphs]
    end = next(i for i, t in enumerate(src) if t.startswith("第４章"))
    doc = Document()
    out = []            # (text, 段落内改行するか)
    pending_title = None
    art_n = 0
    for t in src[:end]:
        if not t.strip():
            out += [("", False)] * (1 + art_n % 2)       # 空段落が1〜2個
            continue
        m = re.match(r"^第(.+?)章[　 ]+(.+)$", t)
        if m:
            n, title = m.groups()
            variants = [f"第{n}章　{title}",
                        f"第 {half(n)} 章　　{'　'.join(title)}",
                        f"　　第{half(n)}章 {title}  "]
            out.append((variants[len(out) % 3], False))
            continue
        m = re.match(r"^（(.+)）$", t)
        if m:
            pending_title = m.group(1)
            continue
        m = re.match(r"^第(.+?)条[　 ](.*)$", t)
        if m:
            art_n += 1
            n, body = m.groups()
            title = pending_title or ""
            pending_title = None
            k = art_n % 5
            if k == 0:      # 見出しと本文が同じ行、半角番号・半角括弧
                out.append((f"第{half(n)}条({title}) {body}", False))
            elif k == 1:    # 見出しを別行に、括弧内にスペース
                out.append((f"（ {' '.join(title)} ）", False))
                out.append((f"第 {half(n)} 条　{body}", False))
            elif k == 2:    # 番号の後ろに見出し、本文は次の段落
                out.append((f"第{n}条　　（{title}）", False))
                out.append((f"　{body}", False))
            elif k == 3:    # 文の途中で段落を分けた見た目の改行
                out.append((f"（{title}）", False))
                cut = max(len(body) // 2, 1)
                out.append((f"第{n} 条　{body[:cut]}", False))
                out.append((body[cut:] + "　", False))
            else:           # 段落内改行
                out.append((f"（{title}）", False))
                out.append((f"第{n}条 {body}", True))
            continue
        m = re.match(r"^([２-９])[　 ](.*)$", t)
        if m:           # 項番号
            n, body = m.groups()
            forms = [f"{n}　{body}", f"{half(n)}. {body}",
                     f"{half(n)} .　{body}", f"　{n}　{body}"]
            out.append((forms[len(out) % 4], False))
            continue
        m = re.match(r"^([①-⑩])[　 ](.*)$", t)
        if m:           # 号番号
            c, body = m.groups()
            i = CIRCLED.index(c) + 1
            forms = [f"{c}　{body}", f"({i}) {body}", f"　　{i})　{body}"]
            out.append((forms[art_n % 3], False))
            continue
        out.append((t + " ", False))
    for text, soft in out:
        p = doc.add_paragraph()
        if soft and "、" in text:
            a, b = text.split("、", 1)
            p.add_run(a + "、").add_break()
            p.add_run(b)
        else:
            p.add_run(text)
    doc.save(DST)
    print(DST)


if __name__ == "__main__":
    main()
