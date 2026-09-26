# -*- coding: utf-8 -*-
"""
就業規則サンプル（tests/samples/*.docx）による回帰テスト。

各サンプルを解析・整形（テンプレ t1〜t4）し、結果を tests/expected/ の
基準と比べる。整形プログラムを直したあとに実行し、既存の結果が
変わっていないか（壊れていないか）を確かめるためのもの。

  python3 tests/regression.py            # 基準と比較
  python3 tests/regression.py --update   # 意図した変更のあと基準を作り直す

比較するもの（サンプルの記述内容そのものは検証しない）
  <名前>.parsed.json     : 解析結果（章・条の構造）
  <名前>.<テンプレ>.txt  : 整形後 docx の段落テキストとスタイル名
"""

import contextlib
import difflib
import io
import json
import os
import shutil
import sys
import tempfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "py"))

from docx import Document  # noqa: E402
from kisoku_parser import parse_kisoku  # noqa: E402
from apply_style import apply_style  # noqa: E402

SAMPLES = os.path.join(ROOT, "tests", "samples")
EXPECTED = os.path.join(ROOT, "tests", "expected")
TEMPLATES = {f"t{i}": os.path.join(ROOT, "templates", f"就業規則テンプレ{i}.docx")
             for i in range(1, 5)}


def _docx_text(path):
    doc = Document(path)
    lines = [f"[{p.style.name}] {p.text}" for p in doc.paragraphs]
    lines += [f"[表{i}] " + " | ".join(c.text for c in row.cells)
              for i, t in enumerate(doc.tables) for row in t.rows]
    return "\n".join(lines) + "\n"


def run_sample(src, work):
    """1ファイル分の {基準ファイル名: 内容} を返す。"""
    name = os.path.splitext(os.path.basename(src))[0]
    local = os.path.join(work, os.path.basename(src))
    shutil.copy(src, local)
    out = {}
    with contextlib.redirect_stdout(io.StringIO()):
        tree = parse_kisoku(local, use_gemini=False)
    json_path = os.path.join(work, name + "_parsed.json")
    text = json.dumps(tree, ensure_ascii=False, indent=2) + "\n"
    with open(json_path, "w", encoding="utf-8") as f:
        f.write(text)
    out[f"{name}.parsed.json"] = text
    for key, tpl in TEMPLATES.items():
        dst = os.path.join(work, f"{name}_{key}.docx")
        with contextlib.redirect_stdout(io.StringIO()):
            apply_style(json_path=json_path, template_path=tpl,
                        template_key=key, output_path=dst,
                        source_docx=local, show_notice=False)
        out[f"{name}.{key}.txt"] = _docx_text(dst)
    return out


def main():
    update = "--update" in sys.argv
    os.makedirs(EXPECTED, exist_ok=True)
    samples = sorted(f for f in os.listdir(SAMPLES) if f.endswith(".docx"))
    failed = []
    for fn in samples:
        with tempfile.TemporaryDirectory() as work:
            try:
                results = run_sample(os.path.join(SAMPLES, fn), work)
            except Exception as e:  # 例外で落ちるのも「壊れた」扱い
                print(f"NG  {fn}: 例外 {type(e).__name__}: {e}")
                failed.append(fn)
                continue
        for exp_name, actual in results.items():
            exp_path = os.path.join(EXPECTED, exp_name)
            if update:
                with open(exp_path, "w", encoding="utf-8") as f:
                    f.write(actual)
                continue
            if not os.path.exists(exp_path):
                print(f"NG  {exp_name}: 基準がありません（--update で作成）")
                failed.append(exp_name)
                continue
            with open(exp_path, encoding="utf-8") as f:
                expected = f.read()
            if expected != actual:
                diff = list(difflib.unified_diff(
                    expected.splitlines(), actual.splitlines(),
                    "expected", "actual", lineterm="", n=1))
                print(f"NG  {exp_name}: 基準と差分あり（先頭40行）")
                print("\n".join(diff[:40]))
                failed.append(exp_name)
        if not any(n.startswith(os.path.splitext(fn)[0] + ".") for n in failed):
            print(f"{'UPD' if update else 'OK '} {fn}")
    if failed:
        print(f"\n失敗 {len(failed)} 件")
        sys.exit(1)
    print("\n基準を更新しました" if update else "\nすべて基準どおりです")


if __name__ == "__main__":
    main()
