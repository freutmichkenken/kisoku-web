# -*- coding: utf-8 -*-
"""custom_template の最小検証。 python3 py/test_custom_template.py で実行。"""

import contextlib
import copy
import io
import json
import os
import re
import tempfile
import zipfile

import custom_template as ct
from apply_style import apply_style

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BASES = {f"t{i}": os.path.join(ROOT, "templates", f"就業規則テンプレ{i}.docx")
         for i in range(1, 5)}


def _lv(fmt, pre, suf, sep, first, left, size=10.5, bold=False,
        align="left", before=0, after=0):
    return {"fmt": fmt, "pre": pre, "suf": suf, "sep": sep, "first": first,
            "left": left, "size": size, "bold": bold, "font": "",
            "align": align, "before": before, "after": after}


SETTINGS = {
    "version": 1, "base": "t1", "layout": "inline", "font": "ＭＳ 明朝",
    "size": 10.5, "line": 16, "table_size": 10.5,
    "levels": {
        "chapter": _lv("decimalFullWidth", "第", "章", "zen", 0, 0, 16, True,
                       "center", 18, 6),
        "article": _lv("decimalFullWidth", "第", "条", "space", 0, 0, 12,
                       before=12),
        "para1": _lv("decimal", "", "", "tab", 0, 2),
        "paraN": _lv("decimal", "", "", "tab", 0, 2),
        "item": _lv("decimal", "(", ")", "tab", 3.5, 5),
        "sub": _lv("decimal", "", ".", "tab", 5.5, 8),
        "sub_single": _lv("none", "……", "", "zen", 6, 9),
        "sub2": _lv("aiueoFullWidth", "(", ")", "tab", 7.5, 10),
    },
}

TREE = [{"type": "chapter", "title": "総則", "articles": [
    {"type": "article", "title": "目的", "paragraphs": [
        {"number": 1, "body": "この規則は、従業員の労働条件を定める。", "items": [
            {"body": "号の一", "sub_items": [
                {"body": "下位の一", "sub_items2": [{"body": "ア"}, {"body": "イ"}]},
                {"body": "下位の二"}]},
            {"body": "号の二"}]},
        {"number": 2, "body": "前項に定めのない事項は法令による。", "items": []}]},
]}, {"type": "chapter", "title": "人事", "articles": [
    {"type": "article", "title": "採用", "paragraphs": [
        {"number": 1, "body": "会社は採用する。", "items": []}]}]}]


def _xml(path, part):
    with zipfile.ZipFile(path) as z:
        return z.read(part).decode("utf-8")


def _lvl_texts(numbering_xml):
    """最後に追加した9段の番号定義の lvlText を段順に返す。"""
    a = re.findall(r'<w:abstractNum [^>]*>(?:(?!</w:abstractNum>).)*'
                   r'<w:nsid w:val="4B570001"/>.*?</w:abstractNum>',
                   numbering_xml, re.S)[-1]
    return re.findall(r'<w:lvlText w:val="([^"]*)"/>', a)


def _run(settings, work, insert_toc=True):
    tpl = os.path.join(work, "tpl.docx")
    M = ct.build(settings, BASES, tpl)
    jp = os.path.join(work, "tree.json")
    with open(jp, "w", encoding="utf-8") as f:
        json.dump(TREE, f, ensure_ascii=False)
    out = os.path.join(work, "out.docx")
    with contextlib.redirect_stdout(io.StringIO()):
        apply_style(json_path=jp, template_path=tpl, template_key="custom",
                    template_map=M, output_path=out, show_notice=False,
                    keep_crossref=False, verify=False, insert_toc=insert_toc)
    return M, out


with tempfile.TemporaryDirectory() as work:
    # --- 1行型：条見出しに条番号、項は 1, 2 ---
    M, out = _run(SETTINGS, work)
    num = _xml(out, "word/numbering.xml")
    texts = _lvl_texts(num)
    assert texts == ["", "第%2章　", "第%3条", "%4", "%5", "(%6)", "%7.",
                     "……\u3000", "(%9)"], texts
    sty = _xml(out, "word/styles.xml")
    for sid in ("KW-1", "KW-2", "KW-3", "KW-4", "KW-5", "KW-6", "KW-7",
                "KW-8", "KW-9"):
        assert f'w:styleId="{sid}"' in sty, sid
    # 1字 = 10.5pt = 210twips。号は2行目5字・1行目3.5字 → ぶら下げ1.5字
    assert re.search(r'w:styleId="KW-5".*?<w:ind w:left="1050" '
                     r'w:hanging="315"/>', sty, re.S)
    # 章は16ptなので行の高さは 16*1.25=20pt（400twips）に広がる
    assert re.search(r'w:styleId="KW-1".*?w:line="400"', sty, re.S)
    # 条は章をまたいで通し番号（lvlRestart=0）、第2項以降は条見出しで振り直す
    assert re.search(r'<w:lvl w:ilvl="2">(?:(?!</w:lvl>).)*'
                     r'<w:lvlRestart w:val="0"/>', num, re.S)
    assert re.search(r'<w:lvl w:ilvl="4"><w:start w:val="2"/>'
                     r'(?:(?!</w:lvl>).)*<w:lvlRestart w:val="3"/>', num, re.S)
    # 1行型では目次スタイル toc 2 に番号を付けない（テンプレ1は元から無い）
    assert not re.search(r'<w:name w:val="toc 2"/>(?:(?!</w:style>).)*numPr',
                         sty, re.S)
    doc = _xml(out, "word/document.xml")
    assert 'TOC \\h \\z \\t "KW-1 章,1,KW-2 条見出し,2"' in doc, "目次"
    # 設定が埋め込まれ、読み戻せる
    assert ct.read_settings(out) == ct.normalize(SETTINGS)

    # --- 分離型：見出し行は番号なし、第1項に条番号。土台がテンプレ1でも
    #     目次 toc 2 に条番号の連番が付く ---
    sep = copy.deepcopy(SETTINGS)
    sep["layout"] = "separate"
    sep["levels"]["article"]["sep"] = "zen"
    sep["levels"]["para1"]["fmt"] = "none"   # 分離型では使わない
    M, out = _run(sep, work)
    texts = _lvl_texts(_xml(out, "word/numbering.xml"))
    assert texts[2] == "" and texts[3] == "第%4条　", texts
    sty = _xml(out, "word/styles.xml")
    assert re.search(r'<w:name w:val="toc 2"/>(?:(?!</w:style>).)*numPr',
                     sty, re.S), "toc 2 に番号が無い"

    # --- テンプレ2を土台に1行型 → 元の toc 2 の番号は外す ---
    t2 = copy.deepcopy(SETTINGS)
    t2["base"] = "t2"
    M, out = _run(t2, work)
    sty = _xml(out, "word/styles.xml")
    assert not re.search(r'<w:name w:val="toc 2"/>(?:(?!</w:style>).)*numPr',
                         sty, re.S), "toc 2 の番号が残っている"
    # docVars は settings.xml のスキーマ順（compat の後、rsids の前）
    st = _xml(out, "word/settings.xml")
    assert st.index("<w:compat") < st.index("<w:docVars") < st.index("<w:rsids")

    # 号の下位その2 は独自の字下げ（1行目6字・2行目以降9字）
    assert re.search(r'w:styleId="KW-7".*?<w:ind w:left="1890" '
                     r'w:hanging="630"/>', sty, re.S)

# --- 整形済み docx から見た目どおりの設定を読み取る（settings_reader） ---
import settings_reader as sr
from docx import Document
from docx.oxml.ns import qn


def _template_out(key, work):
    jp = os.path.join(work, "tree.json")
    with open(jp, "w", encoding="utf-8") as f:
        json.dump(TREE, f, ensure_ascii=False)
    out = os.path.join(work, f"out_{key}.docx")
    with contextlib.redirect_stdout(io.StringIO()):
        apply_style(json_path=jp, template_path=BASES[key], template_key=key,
                    output_path=out, show_notice=False, keep_crossref=False,
                    verify=False)
    return out


with tempfile.TemporaryDirectory() as work:
    # カスタムで整形した docx は、設定がそのまま読み戻せる（1行型・分離型）
    for s in (SETTINGS, sep):
        _M, out = _run(s, work)
        r = sr.read_from_docx(out)
        assert r["settings"] == ct.normalize(s), r
        assert r["notes"] == [], r["notes"]

    # Word での手直し：スタイルの変更（号を12pt・番号定義の字下げ）と、
    # 段落を直接変更（号を中央揃え）のどちらも見た目どおりに読む
    _M, out = _run(SETTINGS, work)
    d = Document(out)
    st = next(x for x in d.styles.element.findall(qn("w:style"))
              if x.get(qn("w:styleId")) == "KW-5")
    st.find(qn("w:rPr") + "/" + qn("w:sz")).set(qn("w:val"), "24")
    D = sr._Doc(d)
    ind = D.lvl(_M["item"][1], 5).find(qn("w:pPr") + "/" + qn("w:ind"))
    ind.set(qn("w:left"), "1260")          # 2行目以降 6字
    for p in d.paragraphs:
        if p.style.style_id == "KW-5":
            p.paragraph_format.alignment = 1   # 中央揃え
    edited = os.path.join(work, "edited.docx")
    d.save(edited)
    lv = sr.read_from_docx(edited)["settings"]["levels"]["item"]
    assert (lv["size"], lv["left"], lv["first"], lv["align"]) \
        == (12, 6, 4.5, "center"), lv

    # テンプレ1〜4で整形した docx：どのテンプレかを見分け、番号の形を読む
    for key, layout in (("t1", "inline"), ("t2", "separate"),
                        ("t3", "separate"), ("t4", "inline")):
        r = sr.read_from_docx(_template_out(key, work))
        s = r["settings"]
        assert (s["base"], s["layout"]) == (key, layout), (key, s)
        art = s["levels"]["article"]
        assert (art["pre"], art["suf"]) == ("第", "条"), (key, art)
        it = s["levels"]["item"]
        assert (it["fmt"], it["pre"], it["suf"]) == ("decimal", "(", ")"), it
        assert s["levels"]["sub_single"]["pre"] == "……", key
        assert "para1" not in s["levels"] or layout == "inline" \
            or "fmt" not in s["levels"]["para1"], key

# このアプリで整形していない docx（整形前の就業規則）は読み取れない
try:
    sr.read_from_docx(os.path.join(ROOT, "tests", "samples",
                                   "messy_kisoku.docx"))
except ValueError:
    pass
else:
    raise AssertionError("整形していない docx を読み取った")

    # --- 丸数字の号が21個以上 → ㉑以降は番号を文字で書き、自動番号を外す ---
    circ = copy.deepcopy(SETTINGS)
    circ["levels"]["item"]["fmt"] = "decimalEnclosedCircle"
    circ["levels"]["item"]["pre"] = ""
    circ["levels"]["item"]["suf"] = ""
    circ["levels"]["item"]["sep"] = "zen"
    TREE_BAK = copy.deepcopy(TREE)
    TREE[0]["articles"][0]["paragraphs"][0]["items"] = [
        {"body": f"項目{i}"} for i in range(1, 53)]
    M, out = _run(circ, work)
    doc = _xml(out, "word/document.xml")
    assert "㉑\u3000項目21" in doc and "㊿\u3000項目50" in doc, "㉑〜㊿"
    assert "51\u3000項目51" in doc, "51以降は普通の数字"
    assert "⑳" not in doc, "⑳までは自動番号のまま"
    assert "丸数字（①②…）の番号は" in _xml(out, "word/comments.xml")
    TREE[:] = TREE_BAK


# 番号定義の読み取り：上の段の番号を含む形式は読めない、全角スペースは sep
from docx.oxml import parse_xml as _px
_W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
_lv = lambda fmt, text: _px(  # noqa: E731
    f'<w:lvl {_W} w:ilvl="5"><w:numFmt w:val="{fmt}"/>'
    f'<w:lvlText w:val="{text}"/></w:lvl>')
assert sr._numbering(_lv("decimal", "%5-%6"))[0] is None
assert sr._numbering(_lv("decimal", "（%6）　"))[0] == {
    "fmt": "decimal", "pre": "（", "suf": "）", "sep": "zen"}
assert sr._numbering(_lv("bullet", "・"))[0] is None
# Word で番号を「なし」にした段落（numId=0）は番号なし
assert sr._numbering(None, num_id="0")[0] == {"fmt": "none"}
assert sr._numbering(None, num_id="99")[0] is None

# --- 号の下位その2 が無い古い設定は、その1と同じ字下げ・記号なしで補う ---
old = copy.deepcopy(SETTINGS)
del old["levels"]["sub_single"]
n = ct.normalize(old)["levels"]["sub_single"]
assert n["pre"] == "" and n["left"] == 8 and n["first"] == 5.5, n

# --- 不正な入力は ValueError ---
for path, value in ((("levels", "item", "first"), -1),
                    (("levels", "item", "pre"), "%"),
                    (("levels", "item", "fmt"), "bullet"),
                    (("size",), "10"),
                    (("base",), "t9"),
                    (("levels", "item", "pre"), "あ" * 9)):
    bad = copy.deepcopy(SETTINGS)
    d = bad
    for k in path[:-1]:
        d = d[k]
    d[path[-1]] = value
    try:
        ct.normalize(bad)
    except ValueError:
        pass
    else:
        raise AssertionError(f"検証をすり抜けた: {path}={value!r}")
bad = copy.deepcopy(SETTINGS)
del bad["levels"]["sub2"]
try:
    ct.normalize(bad)
except ValueError:
    pass
else:
    raise AssertionError("段が欠けた設定を通した")

bad = copy.deepcopy(SETTINGS)
bad["version"] = True
try:
    ct.normalize(bad)
except ValueError:
    pass
else:
    raise AssertionError("version=true を通した")

# --- docVars は末尾の拡張要素より前、スキーマ上の後続要素の前に入る ---
from docx import Document
from docx.oxml import parse_xml
_d = Document(BASES["t1"])
_sel = _d.settings.element
for _c in list(_sel):
    if _c.tag.split("}")[-1] in ct._AFTER_DOCVARS:
        _sel.remove(_c)
_sel.append(parse_xml('<w14:docId xmlns:w14="http://schemas.microsoft.com/'
                      'office/word/2010/wordml" w14:val="1"/>'))
ct.write_settings(_d, {"version": 1})
_tags = [c.tag.split("}")[-1] for c in _sel]
assert _tags.index("compat") < _tags.index("docVars") < _tags.index("docId"), _tags

# --- 設定の入っていない docx は読み込めない ---
try:
    ct.read_settings(BASES["t1"])
except ValueError:
    pass
else:
    raise AssertionError("設定の無い docx を読み込んだ")
print("OK")
