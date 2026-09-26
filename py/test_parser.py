# -*- coding: utf-8 -*-
"""kisoku_parser の最小検証。 python3 py/test_parser.py で実行。"""

from kisoku_parser import classify_line, merge_lines

# 番号の前後に空白がある「第 24 条 (…)」も条見出しとして分けて認識する
_m = merge_lines([('text', '…できるものとする。'), ('text', '第 35 条 (休職期間)')])
assert len(_m) == 2, _m
_n = classify_line('第 24 条 (制裁の種類、程度)')
assert (_n['type'], _n['number_str'], _n['title']) == ('article', '第24条', '制裁の種類、程度'), _n
# 行頭の条文参照「第 9 条に定める…」は条見出しにしない
assert classify_line('第 9 条に定める手続による。')['type'] != 'article'
print("OK")

# Gemini判定は本文扱いの段落で呼ばれ、章・条の判定は採用しない
import gemini_helper, kisoku_parser, contextlib, io, os
from docx import Document
_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), '_test_gemini.docx')
_d = Document()
_d.add_paragraph('第1条 (目的)')
_d.add_table(rows=1, cols=1)     # 表の直後の行は前の行と結合されない
_d.add_paragraph('【火曜日】')
_d.save(_path)
_calls = []
gemini_helper.is_available = lambda: True
gemini_helper.classify_with_gemini = lambda t, b, a: (_calls.append(t), 'article')[1]
kisoku_parser.apply_gemini_split = lambda *a, **k: None
try:
    with contextlib.redirect_stdout(io.StringIO()):
        _tree = kisoku_parser.parse_kisoku(_path, use_gemini=True)
finally:
    os.remove(_path)
assert _calls == ['【火曜日】'], _calls
assert sum(1 for n in _tree if n['type'] == 'article') == 1, _tree
print("OK (gemini)")
